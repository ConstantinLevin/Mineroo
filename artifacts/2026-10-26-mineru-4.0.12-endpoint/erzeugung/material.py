"""Baut das Testmaterial der Aufnahme, erfunden und ohne personenbezogene Daten.

Aufruf: material.py ZIELORDNER. Erzeugt:
  normale_seite.pdf   eine Seite mit Kopfzeile, Fusszeile, Seitenzahl, Randtext,
                      Seitenfussnote, Fliesstext, Bild und Tabelle
  leere_seite.png     eine leere A4-Seite als Bilddatei (150 dpi)
  mehrteilig.pdf      zwoelf Seiten mit Kopf-/Fusszeile, Text und je einem Bild
  dokument.docx       ein Word-Dokument mit Ueberschrift, Text und Tabelle
  seite.html          eine HTML-Seite mit Ueberschrift, Text, Liste und Tabelle
  kaputte_pixel.png   ein PNG mit intaktem Kopf und kaputten Pixeldaten: der
                      zlib-Strom im IDAT nach seinem Kopf mit 0xFF ueberschrieben,
                      die Pruefsumme des Abschnitts neu gerechnet; Pillow oeffnet
                      es und scheitert erst beim Dekodieren der Pixel
"""
import io
import pathlib
import struct
import sys
import zlib

from docx import Document
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

ziel = pathlib.Path(sys.argv[1])
ziel.mkdir(parents=True, exist_ok=True)
breite, hoehe = A4
FARBEN = ["#c0392b", "#2980b9", "#27ae60", "#8e44ad", "#d35400", "#16a085", "#2c3e50", "#f39c12"]


def diagramm(nummer: int) -> ImageReader:
    bild = Image.new("RGB", (480, 280), "white")
    zeichner = ImageDraw.Draw(bild)
    zeichner.line([(30, 250), (460, 250)], fill="black", width=2)
    zeichner.line([(30, 20), (30, 250)], fill="black", width=2)
    for q in range(4):
        h = 40 + ((nummer + q) * 37) % 180
        zeichner.rectangle([60 + q * 100, 250 - h, 130 + q * 100, 250], fill=FARBEN[(nummer + q) % len(FARBEN)])
        zeichner.text((80 + q * 100, 255), f"Q{q + 1}", fill="black")
    puffer = io.BytesIO()
    bild.save(puffer, format="PNG")
    return ImageReader(io.BytesIO(puffer.getvalue()))


def tabelle(c: canvas.Canvas, x: float, y: float, reihen: list[tuple[str, ...]]) -> None:
    c.setFont("Helvetica", 10)
    for r, reihe in enumerate(reihen):
        for s, wert in enumerate(reihe):
            c.rect(x + s * 110, y - r * 20, 110, 20)
            c.drawString(x + 6 + s * 110, y - r * 20 + 6, wert)


# 1. Die normale Seite mit allen fuenf Randklassen.
c = canvas.Canvas(str(ziel / "normale_seite.pdf"), pagesize=A4)
c.setTitle("Jahresbericht 2025 der Beispielwerke")
c.setFont("Helvetica", 9)
c.drawString(50, hoehe - 36, "Beispielwerke AG  |  Jahresbericht 2025  |  Lagebericht")
c.line(50, hoehe - 42, breite - 50, hoehe - 42)
c.setFont("Helvetica-Bold", 18)
c.drawString(50, hoehe - 90, "3. Geschaeftsverlauf im Berichtsjahr")
c.setFont("Helvetica", 11)
absatz = [
    "Die Beispielwerke AG steigerte ihren Umsatz im Geschaeftsjahr 2025 auf 412,7 Millionen Euro",
    "und lag damit 6,3 Prozent ueber dem Vorjahr. Getragen wurde das Wachstum vor allem vom",
    "Segment Anlagenbau, dessen Auftragseingang sich auf 158,2 Millionen Euro erhoehte.¹",
    "Das operative Ergebnis (EBIT) erreichte 37,9 Millionen Euro; die EBIT-Marge lag bei 9,2 Prozent.",
    "Die Investitionen stiegen auf 21,4 Millionen Euro und flossen ueberwiegend in neue Fertigungslinien.",
]
y = hoehe - 120
for zeile in absatz:
    c.drawString(50, y, zeile)
    y -= 16
c.drawImage(diagramm(1), 50, y - 200, width=300, height=175)
c.setFont("Helvetica-Oblique", 10)
c.drawString(50, y - 215, "Abbildung 3: Umsatz je Quartal 2025 in Millionen Euro")
c.setFont("Helvetica-Bold", 11)
c.drawString(50, y - 245, "Tabelle 2: Kennzahlen im Ueberblick")
tabelle(c, 50, y - 275, [("Kennzahl", "2024", "2025"), ("Umsatz (Mio. EUR)", "388,2", "412,7"),
                         ("EBIT (Mio. EUR)", "33,1", "37,9"), ("Mitarbeitende", "1.904", "1.988")])
c.setFont("Helvetica", 8)
c.line(50, 92, 200, 92)
c.drawString(50, 80, "¹ Auftragseingang einschliesslich Rahmenvertraegen mit einer Laufzeit ueber zwoelf Monate.")
c.setFont("Helvetica", 9)
c.drawString(50, 40, "Beispielwerke AG, Musterstrasse 1, 12345 Beispielstadt")
c.drawRightString(breite - 50, 40, "Seite 14")
c.saveState()
c.translate(breite - 22, 300)
c.rotate(90)
c.setFont("Helvetica", 8)
c.drawString(0, 0, "Geprueft durch die Revisionsgesellschaft Beispiel GmbH – Fassung vom 12.03.2026")
c.restoreState()
c.showPage()
c.save()

# 2. Die leere Seite als Bilddatei.
Image.new("RGB", (1240, 1754), "white").save(ziel / "leere_seite.png")

# 3. Zwoelf Seiten fuer den Schnitt in drei Teilstuecke.
c = canvas.Canvas(str(ziel / "mehrteilig.pdf"), pagesize=A4)
c.setTitle("Technischer Bericht Kuehlanlagen")
for seite in range(12):
    c.setFont("Helvetica", 9)
    c.drawString(50, hoehe - 36, "Technischer Bericht Kuehlanlagen – Messkampagne 2025")
    c.drawRightString(breite - 50, 36, f"Seite {seite + 1} von 12")
    c.setFont("Helvetica-Bold", 15)
    c.drawString(50, hoehe - 85, f"{seite + 1}. Anlage {chr(65 + seite)}: Messwerte und Bewertung")
    c.setFont("Helvetica", 11)
    zeilen = [
        f"An Anlage {chr(65 + seite)} wurde ueber vier Quartale die Leistungszahl gemessen.",
        f"Der Mittelwert lag bei {3 + seite * 0.1:.1f}, die Spitzenlast bei {120 + seite * 7} Kilowatt.",
        "Die Abbildung zeigt die quartalsweise Entwicklung der elektrischen Aufnahme.",
    ]
    y = hoehe - 115
    for zeile in zeilen:
        c.drawString(50, y, zeile)
        y -= 16
    c.drawImage(diagramm(seite), 50, y - 190, width=300, height=175)
    c.setFont("Helvetica-Oblique", 10)
    c.drawString(50, y - 205, f"Abbildung {seite + 1}: Aufnahme der Anlage {chr(65 + seite)} je Quartal")
    c.showPage()
c.save()

# 4. Ein Word-Dokument.
dokument = Document()
dokument.add_heading("Protokoll der Werkstattbesprechung", level=1)
dokument.add_paragraph("Die Besprechung behandelte die Wartungsplanung fuer das zweite Halbjahr.")
dokument.add_heading("Beschluesse", level=2)
dokument.add_paragraph("Die Pruefung der Druckluftanlage wird auf September vorgezogen.", style="List Bullet")
dokument.add_paragraph("Fuer die Lackiererei werden zwei neue Filtersaetze beschafft.", style="List Bullet")
t = dokument.add_table(rows=3, cols=3)
for r, reihe in enumerate([("Anlage", "Termin", "Dauer"), ("Druckluft", "09/2026", "2 Tage"), ("Lackiererei", "10/2026", "1 Tag")]):
    for s, wert in enumerate(reihe):
        t.cell(r, s).text = wert
dokument.save(ziel / "dokument.docx")

# 5. Eine HTML-Seite.
(ziel / "seite.html").write_text(
    """<!DOCTYPE html>
<html lang="de"><head><meta charset="utf-8"><title>Hausmitteilung Energie</title></head>
<body>
<h1>Hausmitteilung: Energiesparen im Winter</h1>
<p>Ab dem 1. November gelten in allen Gebaeuden neue Heizzeiten. Bitte beachten Sie die folgenden Punkte.</p>
<ul><li>Raumtemperatur in Bueros: 20 Grad</li><li>Flure und Lager: 15 Grad</li></ul>
<table><tr><th>Gebaeude</th><th>Heizbeginn</th><th>Heizende</th></tr>
<tr><td>Halle 1</td><td>06:00</td><td>18:00</td></tr><tr><td>Verwaltung</td><td>07:00</td><td>17:30</td></tr></table>
</body></html>
""",
    encoding="utf-8",
)
# 6. Ein PNG, dessen Kopf stimmt und dessen Pixeldaten kaputt sind.
bild = Image.new("RGB", (800, 500), "white")
zeichner = ImageDraw.Draw(bild)
for i in range(20):
    zeichner.text((40, 20 + i * 22), f"Zeile {i + 1}: Prueftext fuer die Bilddekodierung", fill="black")
puffer = io.BytesIO()
bild.save(puffer, format="PNG")
daten = bytearray(puffer.getvalue())
pos = 8
while pos < len(daten):
    laenge = struct.unpack(">I", daten[pos : pos + 4])[0]
    if daten[pos + 4 : pos + 8] == b"IDAT":
        anfang = pos + 8
        daten[anfang + 2 : anfang + laenge] = b"\xff" * (laenge - 2)
        crc = zlib.crc32(bytes(daten[pos + 4 : anfang + laenge])) & 0xFFFFFFFF
        daten[anfang + laenge : anfang + laenge + 4] = struct.pack(">I", crc)
    pos += 12 + laenge
(ziel / "kaputte_pixel.png").write_bytes(bytes(daten))

print("geschrieben:", sorted(p.name for p in ziel.iterdir()))
